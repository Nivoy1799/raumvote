"""Build PowerPoint deck from rendered slide images + presenter notes.

Run after `pdftoppm -r 200 -png presentation-ikts.pdf slide`.
Output: presentation-ikts.pptx (16:9, image per slide, German speaker notes).
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Emu

HERE = Path(__file__).parent
OUT = HERE / "presentation-ikts.pptx"

# 16:9 widescreen, matches Typst 254 x 143 mm aspect
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NOTES = {
    1: """\
TITELFOLIE — Felix (Scrum Master), ca. 20 Sekunden.

Kontext: Wir sind die ERSTE Gruppe in der Reihenfolge. Kein "wie ihr vorhin gesehen habt" möglich — sauber einsteigen.

Begrüssung:
- "Guten Tag. Wir machen den Anfang."
- "Wir sind Felix, Jovin und Luis und stellen euch unseren Case vor: Beteiligung von Jugendlichen, Produkt RaumVote."
- "Unsere Botschaft: ‚Deine Stimme für deinen Raum.'"

Wichtig kurz erwähnen:
- Wir sind ein interdisziplinäres Team: Felix und Luis aus dem Wirtschaftsingenieurwesen, Jovin aus der Informatik.
- Jovin ist heute nicht vor Ort — sein Beitrag kommt gleich als 1.5-min Video.

Ablauf:
- Rund 7 Minuten Präsentation
- 5 Minuten Live-Voting — ihr macht alle mit
- 5 Minuten Fragen

Übergang: "Beginnen wir mit dem, was uns auf diese Idee gebracht hat."
""",
    2: """\
FREITAGABEND, DORFPLATZ — Felix, ca. 50 Sekunden.

Bildlich einsteigen, nicht ablesen:
- "Stellt euch einen Freitagabend auf einem Dorfplatz vor."
- "Jugendliche treffen sich, hören Musik, hängen ab — für sie ein zentraler sozialer Treffpunkt."
- "Anwohnende beschweren sich über Lärm. Die Gemeinde reagiert — meistens, ohne mit den Jugendlichen zu reden."

Daraus das Muster ableiten:
- Klassische Formate (Vernehmlassungen, Abendsitzungen, formale Eingaben) passen nicht in den Alltag von 14- bis 18-Jährigen.
- Wer keinen Vereinszugang hat — und das ist die grosse Mehrheit — bleibt im Prozess unsichtbar.
- Ohne Vertrauen in den Umgang mit ihren Daten beteiligen sich Jugendliche nicht ehrlich.

Rechte Box — UN-Kinderrechtskonvention Artikel 12:
- "International haben Jugendliche das Recht, in allen sie betreffenden Angelegenheiten gehört zu werden."
- "Die Schweiz hat das ratifiziert. Kommunal wird es selten konsequent umgesetzt."

Übergang: "Wir haben uns gefragt: wie ändert man das? Daraus ist RaumVote entstanden — mit vier klaren Prinzipien."
""",
    3: """\
UNSERE ANTWORT — Felix, ca. 50 Sekunden.

Wichtig: Diese vier Prinzipien sind nicht aus dem Bauch — sie kommen direkt aus den User Stories der drei Stakeholder.

P1 — Anonym & niederschwellig
"Brief mit QR-Code, ein Scan, kein Account. Auch Jugendliche ohne Vereinsanbindung kommen rein."

P2 — Mobile & spielerisch
"Swipe statt Formular. Eine Frage pro Bildschirm. Bewusst angelehnt an die App-Welt, in der Jugendliche zu Hause sind."

P3 — Datenschutz von Grund auf
"Auch wir als Betreiber sehen keine Klarnamen. Identitäten werden vor der Speicherung verschlüsselt — das ist Designprinzip, nicht nachträgliche Pflichtübung."

P4 — Anschlussfähig für die Gemeinde
"Am Ende kriegt die Gemeinde eine klare Rangliste — direkt verwendbar im nächsten Planungsworkshop. Keine Rohdaten, die niemand interpretieren kann."

Übergang: "Wie sieht das in der App konkret aus?"
""",
    4: """\
SO FUNKTIONIERT RAUMVOTE — Felix, ca. 30 Sekunden.

Drei Schritte, klar und schnell:

1. QR scannen
"Brief im Briefkasten. QR drauf. Scan — die Berechtigung wird lokal auf dem Handy gespeichert. Kein Login, kein Passwort."

2. Swipe durch den Baum
"Zwei Bildoptionen auf dem Bildschirm. Wisch nach links oder rechts. Die nächste Ebene baut sich automatisch auf. So findest du Schritt für Schritt deine Idealvorstellung."

3. Wähle deinen Traumraum
"Am Ende landest du bei einer konkreten Option. Die gibst du als Stimme ab. Likes und Kommentare ermöglichen Austausch. Du siehst live, was andere gewählt haben."

Box unten kurz benennen:
"Für die Gemeinde fällt am Ende eine anonyme Präferenzkarte ab — direkt einsetzbar in Planungsworkshops."

Übergang: "Wie das technisch im Hintergrund läuft, zeigt euch jetzt Jovin im Video."

→ Felix oder Luis startet jetzt das Video.
""",
    5: """\
HINTER DEN KULISSEN — Video, 1.5 min, KEINE LIVE-SPRECHE.

Aktion vor dem Video:
- Separates Videofenster (implementation-jovin.mp4) auf den Beamer ziehen oder Vollbild.
- Lautstärke vorher prüfen.

Video-Inhalt (zur eigenen Orientierung, damit ihr in der Q&A bei Bedarf darauf verweisen könnt):
- 0:00–0:20  Architektur — wie die App im Betrieb läuft
- 0:20–0:55  Wie die KI im Hintergrund die Entscheidungsbäume und Bilder erzeugt
- 0:55–1:25  Wie der Datenschutz technisch sichergestellt wird
- 1:25–1:30  Outro

Nach dem Video:
- Zurück zur Präsi.
- Luis übernimmt: "Was wir aus all dem gelernt haben."
""",
    6: """\
WAS WIR GELERNT HABEN — Luis (PO), ca. 60 Sekunden.

Vier inhaltliche Erkenntnisse:

1. Konzept vor Code
"Wir haben mit Design Thinking begonnen, nicht mit Programmieren. Diese Reihenfolge hat sich bewährt: Erst verstehen, wer was wirklich braucht, dann bauen. So entsteht kein Tool für sich selbst."

2. Drei Stakeholder, ein Tool
"Jugendliche wollen Niederschwelligkeit. Die Jugendarbeit braucht Moderationsrechte gegen Missbrauch. Die Gemeinde will brauchbare Resultate. Diese drei Perspektiven in einem Format auszubalancieren war die Königsdisziplin."

3. Niederschwellig hat Grenzen
"Ein QR-Code kann fotografiert und weitergeleitet werden — theoretisch kann ein Token viral gehen. Wir haben dieses Risiko bewusst akzeptiert und administrative Steuerung eingebaut: verdächtige Tokens können deaktiviert werden."

4. Test mit echten Jugendlichen steht aus
"Im Semester konnten wir das nicht mehr organisieren. Aber: das vollständige Testkonzept liegt im Bericht und kann von der Gemeinde direkt für die Pilotphase verwendet werden."

Persönlich (blau hervorgehoben am unteren Rand) — wichtig, kurz aber ehrlich:
"Persönlich: Wir haben uns als Team sehr gut ergänzt. Die Zusammenarbeit zwischen Wirtschaftsingenieurwesen und Informatik war für uns alle eine echte Bereicherung — fachlich und menschlich."

Übergang: "Genug Theorie. Jetzt seid ihr dran."
""",
    7: """\
DEMO — Luis, Aufforderung ca. 20 Sekunden, dann 5 min Live-Voting.

Sofort-Action:
- "Bitte den QR-Code auf eurem Brief scannen — wir machen jetzt ein Live-Voting."
- Beamer auf RaumVote-App umschalten (Browser mit aktiver Session vorbereitet).

Ablauf (5 Minuten, Felix moderiert mit):
- 0:00–0:30  Alle scannen, landen im Startbildschirm. Beamer zeigt den Einstieg.
- 0:30–3:00  Alle navigieren durch denselben Beispiel-Baum. Kurze Anleitung: "Wisch zu der Option, die euch mehr anspricht."
- 3:00–4:00  Beamer wechselt auf die Resultate — Live-Rangliste.
- 4:00–5:00  Kurze Reflexion mit dem Publikum: "Was hat euch beim Voten überrascht? Hat eine Option gewonnen, die ihr nicht erwartet hättet?"

Backup-Pläne (die Dozierenden empfehlen explizit ein Demo-Video als Backup):
- WLAN-Problem: Hotspot vom Handy bereithalten.
- Komplettausfall: Demo-Backup-Video als mp4 lokal bereit (60–90 Sekunden, zeigt den Flow inkl. Resultate). VOR der Präsi testen.
- Ankündigen statt Verschweigen, falls etwas hängt: "Wir nehmen kurz das Backup-Video — die Idee ist dieselbe."

Übergang: "Bevor wir zu euren Fragen kommen — kurz noch zur KI-Transparenz."
""",
    8: """\
KI-UNTERSTÜTZUNG — Luis oder Felix, ca. 15 Sekunden.

Sachlich und kurz:

"In dieser Arbeit selbst:
- Wir haben Claude und ChatGPT eingesetzt für Struktur, Text und technische Unterstützung beim Code.

Im laufenden RaumVote-Betrieb:
- GPT-4o generiert die Entscheidungsbäume in Echtzeit.
- Gemini erzeugt die zugehörigen Bilder.

KI als Hilfsmittel — alle Inhalte haben wir kritisch geprüft. Die Verantwortung liegt bei uns."

Pflichtfolie an der OST — nicht überspringen.

Übergang: "Wir freuen uns auf eure Fragen."
""",
    9: """\
Q&A — 5 Minuten, Felix (SM) & Luis (PO).

Wer beantwortet was:
- Prozess, Methode, Teamarbeit, Kommunikationsstrategie, Stakeholder → Felix.
- Produkt, User Stories, Beteiligungsablauf, MVP-Scope, Erkenntnisse → Luis.
- Tech-Fragen (Architektur, KI, Datenschutz im Detail) → beide grob aus dem Video heraus. Für Tiefe: "Das beantwortet Jovin gern asynchron — schreibt uns kurz."

Bei Fragen, die im IKTS-Bericht behandelt werden, ruhig darauf verweisen:
- "Im Bericht beschreiben wir das ausführlich — wir können euch das gerne zusenden."
- IKTS-Bericht liegt in docs/documentation/IKTS_Arbeit.pdf.

Wenn das Publikum zögert: erste Frage selber anstossen.
- "Eine Frage, die wir uns auch gestellt haben: Was passiert, wenn ein QR-Code weitergegeben wird?"
- "Was war für euch die schwierigste Designentscheidung?"
- "Wie geht es jetzt weiter mit RaumVote?"

Abschluss: knapper Dank, "Deine Stimme für deinen Raum." als letztes Wort wenn passend.
""",
}


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    blank_layout = prs.slide_layouts[6]  # blank

    for i in range(1, 10):
        img = HERE / f"slide-{i}.png"
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(str(img), 0, 0, width=SLIDE_W, height=SLIDE_H)
        slide.notes_slide.notes_text_frame.text = NOTES[i].strip()

    prs.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size // 1024} KB, {len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
